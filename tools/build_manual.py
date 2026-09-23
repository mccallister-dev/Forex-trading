from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "ORB_ICT_Session_Liquidity_Framework_Manual.docx"

BLUE = RGBColor(46, 116, 181)
DEEP_BLUE = RGBColor(31, 77, 120)
TEAL = RGBColor(0, 128, 128)
ORANGE = RGBColor(224, 124, 35)
INK = RGBColor(32, 42, 52)
MUTED = RGBColor(95, 105, 115)
PALE_BLUE = "E8EEF5"
PALE_TEAL = "E7F3F1"
PALE_ORANGE = "FBEEDF"
LIGHT_GRAY = "F4F6F8"


def set_font(run, size=11, color=INK, bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run(text), 28, DEEP_BLUE, True)
    if subtitle:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.paragraph_format.space_after = Pt(30)
        set_font(s.add_run(subtitle), 14, TEAL, False)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 13 if level == 2 else 12, BLUE if level < 3 else DEEP_BLUE, True)
    return p


def add_para(doc, text, bold_lead=None, after=6, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), size, INK, True)
        set_font(p.add_run(text[len(bold_lead):]), size, INK)
    else:
        set_font(p.add_run(text), size, INK)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), 10.5, INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), 10.5, INK)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(label + "  "), 10.5, DEEP_BLUE, True)
    set_font(p.add_run(text), 10.5, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_header(table.rows[0])
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(value), font_size, DEEP_BLUE, True)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(value), font_size, INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc):
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DEEP_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
    doc.styles["List Bullet"].paragraph_format.left_indent = Inches(0.375)
    doc.styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.188)
    doc.styles["List Number"].paragraph_format.left_indent = Inches(0.375)
    doc.styles["List Number"].paragraph_format.first_line_indent = Inches(-0.188)

    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        add_page_field(footer.paragraphs[0])


def build_manual():
    doc = Document()
    configure_document(doc)

    # Page 1 - cover and first use.
    add_title(doc, "ORB + ICT Session Liquidity Framework", "Five-page quick manual | Pine Script v6")
    add_callout(doc, "Purpose", "A clean, confirmed-bar observation framework for major sessions, one selected opening range, directional validation, and setup sequencing. It does not place trades.", PALE_TEAL)
    add_heading(doc, "First use", 1)
    for step in (
        "Select a 5-minute or 15-minute standard candle chart.",
        "Leave Display mode on Minimal.",
        "Choose the relevant ORB, or leave ORB session on Auto.",
        "Observe Upcoming, Forming, and Closed.",
        "Wait for confirmed validation stages; a score cannot skip the sequence.",
        "Treat outputs as analytical observations, not trade instructions.",
    ):
        add_number(doc, step)
    add_heading(doc, "Install", 2)
    add_para(doc, "Open TradingView > Pine Editor, create a new indicator, replace the sample with the delivered .pine source, and choose Add to chart. The delivered source compiled successfully in TradingView Pine v6 on 20 August 2026.", size=10.5)
    add_callout(doc, "Core rule", "Bias is evidence balance. Setup score is confluence quality. Neither is a probability, win rate, confidence percentage, or recommendation.", PALE_ORANGE)

    doc.add_page_break()

    # Page 2 - modes and settings.
    add_heading(doc, "1. Modes and essential settings", 1)
    add_table(
        doc,
        ["Mode", "What appears"],
        [
            ("Minimal", "Light Asian/London/New York boxes, one selected ORB, close icon, short high/low segments, compact dashboard."),
            ("Standard", "The same clean framework with slightly stronger session and ORB fills."),
            ("Analysis", "Allows opted-in Imbalances, Equal highs/lows, and previous-day/week levels under a strict zone budget."),
            ("Custom", "Allows Advanced technical layers and the optional provisional last-bar label."),
        ],
        [1700, 7660],
    )
    add_heading(doc, "The 40 controls, grouped", 2)
    add_bullet(doc, "General: mode, trading-day retention, display timezone, dashboard, and preferred-session mode/override.")
    add_bullet(doc, "Sessions: show boxes; enable, time, IANA timezone, and color/transparency for Asian, London, and New York.")
    add_bullet(doc, "Opening Range: Auto/Asian/London/New York, duration, box, midpoint, and close marker.")
    add_bullet(doc, "Validation: Standard/Strict, confirmed-bars setting, confirmed markers, and developing labels.")
    add_bullet(doc, "Advanced: optional technical layers, zone budget, alerts, minimum break distance, bias alignment, and disabled-by-default custom session.")
    add_heading(doc, "Session defaults", 2)
    add_table(
        doc,
        ["Session", "Default", "Timezone"],
        [
            ("Asian", "00:00-08:00", "Europe/London"),
            ("London", "08:00-16:00", "Europe/London"),
            ("New York", "08:00-17:00", "America/New_York"),
            ("Custom", "06:00-10:00; off", "Exchange"),
        ],
        [1800, 3000, 4560],
    )
    add_callout(doc, "Timezone note", "Named IANA zones handle daylight-saving changes. Display timezone controls the dashboard clock and retention day key; it does not replace each session's timezone.")

    doc.add_page_break()

    # Page 3 - ORB lifecycle and validation.
    add_heading(doc, "2. ORB lifecycle", 1)
    lifecycle = [
        ("Upcoming", "The selected session's current-day ORB has not started."),
        ("Forming", "High and low update from constituent chart bars."),
        ("Closed", "The final constituent candle confirms; range values freeze permanently."),
        ("Break observed", "A confirmed close clears the frozen range by the minimum distance and passes Strong move qualification."),
        ("Retest", "A later confirmed candle retests the broken boundary and closes on the breakout side."),
        ("Confirmed", "A later confirmed continuation close clears the retest candle's directional extreme."),
    ]
    add_table(doc, ["State", "Meaning"], lifecycle, [2100, 7260], header_fill=PALE_TEAL)
    add_heading(doc, "What freezes and what extends", 2)
    add_bullet(doc, "One diamond is created on the confirmed ORB-completing candle and never moves.")
    add_bullet(doc, "The compact ORB box ends at the ORB deadline.")
    add_bullet(doc, "Only ORB high and low continue after closure, and both stop at session close.")
    add_bullet(doc, "The midpoint is off by default; no session OHLC or full-chart background is drawn.")
    add_heading(doc, "Standard vs Strict", 2)
    add_table(
        doc,
        ["Check", "Standard", "Strict"],
        [
            ("Strong move", "0.75 ATR body; 50% body share", "1.0 ATR body; 60% body share"),
            ("Reversal", "Confirmed reclaim plus direction", "Also clears the preceding bar extreme"),
            ("Retest", "Boundary touch and close beyond", "Also rejection wick or engulfing evidence"),
            ("Continuation", "Clears retest extreme", "Also passes Strong move again"),
        ],
        [1900, 3300, 4160],
        header_fill=PALE_ORANGE,
    )

    doc.add_page_break()

    # Page 4 - bias, preferred session, and score.
    add_heading(doc, "3. Bias, stage, and quality", 1)
    add_para(doc, "Bullish evidence and bearish evidence are accumulated separately. Net bias equals bullish evidence minus bearish evidence.")
    add_table(
        doc,
        ["Net", "Classification", "Evidence sources"],
        [
            ("+5 or more", "Strong Bullish", "Daily open; completed 1H/4H structure; day/Asian sweeps; Imbalance reaction; qualified ORB break"),
            ("+2 to +4", "Bullish", "Same components, weighted independently"),
            ("-1 to +1", "Neutral", "Bullish and bearish evidence broadly balanced"),
            ("-2 to -4", "Bearish", "Same components, weighted independently"),
            ("-5 or less", "Strong Bearish", "Bearish evidence exceeds bullish evidence by five or more"),
        ],
        [1700, 2100, 5560],
        font_size=9,
    )
    add_callout(doc, "Mandatory sequence", "sweep -> reclaim/reversal -> ORB break -> retest/rejection -> continuation", PALE_TEAL)
    add_table(
        doc,
        ["Quality component", "Points"],
        [
            ("Sweep; reclaim/reversal; aligned bias", "2 + 2 + 2"),
            ("ORB break; Strong move; breakout Imbalance", "1 + 2 + 1"),
            ("Retest/rejection; continuation", "2 + 2"),
            ("Target space >= 2R; active selected session", "1 + 1"),
        ],
        [6900, 2460],
        header_fill=LIGHT_GRAY,
    )
    add_para(doc, "Maximum: 16. Labels: Developing 5+, Good 8+, High 11+. A high score never creates a confirmed marker without the full mandatory sequence.", bold_lead="Maximum: ", size=10.5)
    add_heading(doc, "Preferred-session Auto mapping", 2)
    add_bullet(doc, "EUR/GBP/CHF forex -> London; other JPY/AUD/NZD forex -> Asian; crypto -> Asian; everything else -> New York.")
    add_bullet(doc, "Ticker root plus ticker text reduces broker prefix/suffix problems. Manual mode overrides unsupported symbols.")

    doc.add_page_break()

    # Page 5 - retention, alerts, guarantees, limitations.
    add_heading(doc, "4. Retention, alerts, and guarantees", 1)
    add_heading(doc, "Trading-day retention", 2)
    add_para(doc, "Trading days = 3 keeps the current chart day plus the two immediately preceding days that contain bars. Every drawing carries a stable display-timezone day key. Missing market days are skipped. Separate internal safety caps protect TradingView drawing limits.")
    add_heading(doc, "One-shot alerts", 2)
    add_bullet(doc, "ORB closed; confirmed ORB break; retest/rejection confirmed; final setup confirmed.")
    add_bullet(doc, "Dynamic alerts include symbol, session, direction where relevant, setup score, and net bias.")
    add_heading(doc, "Non-repainting checklist", 2)
    for item in (
        "Permanent states, the ORB-close icon, and confirmed setup markers require barstate.isconfirmed.",
        "ORB high/low stop changing immediately after confirmed closure.",
        "Completed 1H/4H values use source offsets; previous-day/week levels use completed periods.",
        "Equal highs/lows use confirmed pivots after right-side bars have elapsed.",
        "No negative history offset, future bar index, or unconfirmed higher-timeframe high/low is used.",
        "Only the optional PROVISIONAL last-bar label may move intrabar.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "Unavoidable limitations", 2)
    add_bullet(doc, "A chart bar is indivisible. A 5-minute ORB on a 15-minute chart uses the first 15-minute candle; use a timeframe no larger than the ORB duration.")
    add_bullet(doc, "Standard candles are required for literal OHLC; synthetic candles change all supplied prices.")
    add_bullet(doc, "Target space is structural geometry, not spread, slippage, execution, or probability.")
    add_callout(doc, "Final reminder", "Observe first, verify the symbol/feed/session alignment, and keep Confirmed bars only enabled for normal use.", PALE_ORANGE)

    doc.core_properties.title = "ORB + ICT Session Liquidity Framework - Quick Manual"
    doc.core_properties.subject = "Five-page installation and observation guide"
    doc.core_properties.keywords = "Pine Script, ORB, sessions, non-repainting, TradingView"
    doc.core_properties.comments = "Generated from the streamlined v6 framework specification."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_manual()
